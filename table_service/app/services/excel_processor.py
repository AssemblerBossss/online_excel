import logging
from collections.abc import Iterable
from datetime import date, datetime, time
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.datetime import to_excel
from openpyxl.workbook import Workbook

from table_service.app.exceptions import EmptyFileException, ValidationException
from table_service.app.schemas.sheet import MAX_COLS, MAX_ROWS

logger = logging.getLogger(__name__)

# Форматы, которые понимает HyperFormula на фронте (dateFormats в engine.ts)
DATE_TEXT_FORMAT = "%d.%m.%Y"
DATETIME_TEXT_FORMAT = "%d.%m.%Y %H:%M:%S"

EXCEL_ERRORS = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}

EXCEL_NUMBER_FORMATS = {
    "number": "#,##0.00",
    "percent": "0.00%",
    "currency": '#,##0.00 "₽"',
    "date": "DD.MM.YYYY",
    "datetime": "DD.MM.YYYY HH:MM:SS",
}


class ExcelProcessorService:
    """Импорт .xlsx в ячейки листа и экспорт обратно. Без доступа к БД."""

    # ---------------------------------------------------------------- импорт

    def parse_workbook(self, file) -> list[dict[str, Any]]:
        """Прочитать первый лист: формулы сохраняются, значения берутся из кэша Excel.

        Возвращает ячейки в формате SheetCell: {row, col, raw, value, style}.
        """
        workbook = load_workbook(file, data_only=False)
        file.seek(0)
        cached_workbook = load_workbook(file, data_only=True)
        sheet = workbook.active
        cached_sheet = cached_workbook.active

        if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLS:
            raise ValidationException(
                f"Лист слишком большой: максимум {MAX_ROWS} строк и {MAX_COLS} колонок"
            )

        cells: list[dict[str, Any]] = []
        for row in sheet.iter_rows():
            for cell in row:
                parsed = self._convert_cell(cell, cached_sheet[cell.coordinate].value)
                if parsed:
                    cells.append(parsed)
        if not cells:
            raise EmptyFileException("Файл не содержит данных для импорта")
        return cells

    def _convert_cell(self, cell, cached_value: Any) -> dict[str, Any] | None:
        content = getattr(
            cell.value, "text", cell.value
        )  # ArrayFormula -> текст формулы
        style = self._style_from_cell(cell)
        raw: str | None = None
        value: Any = None

        if content is None:
            pass
        elif isinstance(content, str) and content.startswith("="):
            raw, value = content, self._json_value(cached_value)
        elif isinstance(content, bool):
            raw, value = ("TRUE" if content else "FALSE"), content
        elif isinstance(content, datetime) and content.time() != time.min:
            raw, value = content.strftime(DATETIME_TEXT_FORMAT), to_excel(content)
            style["numFmt"] = "datetime"
        elif isinstance(content, date):  # datetime ровно в полночь тоже считаем датой
            raw, value = content.strftime(DATE_TEXT_FORMAT), to_excel(content)
            style["numFmt"] = "date"
        elif isinstance(content, int | float):
            raw, value = str(content), content
        else:
            raw, value = str(content), str(content)

        if raw is None and not style:
            return None
        return {
            "row": cell.row - 1,
            "col": cell.column - 1,
            "raw": raw,
            "value": value,
            "style": style or None,
        }

    @staticmethod
    def _json_value(value: Any) -> Any:
        if value is None or isinstance(value, bool | int | float):
            return value
        if isinstance(value, datetime | date):
            return to_excel(value)
        text = str(value)
        return {"e": text} if text in EXCEL_ERRORS else text

    @staticmethod
    def _rgb(color) -> str | None:
        rgb = getattr(color, "rgb", None)
        if isinstance(rgb, str) and len(rgb) == 8:  # AARRGGBB; тема/индексы пропускаем
            return f"#{rgb[2:].lower()}"
        return None

    def _style_from_cell(self, cell) -> dict[str, Any]:
        style: dict[str, Any] = {}
        if cell.fill and cell.fill.fill_type == "solid":
            bg = self._rgb(cell.fill.fgColor)
            if bg and bg != "#ffffff":
                style["bg"] = bg
        if cell.font:
            if cell.font.bold:
                style["bold"] = True
            if cell.font.italic:
                style["italic"] = True
            fg = self._rgb(cell.font.color)
            if fg and fg != "#000000":
                style["fg"] = fg
        if cell.alignment:
            if cell.alignment.horizontal in ("left", "center", "right"):
                style["halign"] = cell.alignment.horizontal
            if cell.alignment.wrap_text:
                style["wrap"] = True
        if cell.number_format and cell.number_format.endswith("%"):
            style["numFmt"] = "percent"
        return style

    # --------------------------------------------------------------- экспорт

    @staticmethod
    def create_streaming_workbook() -> tuple[Workbook, Any]:
        """Начать книгу в write-only режиме."""
        workbook = Workbook(write_only=True)
        return workbook, workbook.create_sheet()

    def append_rows_chunk(
        self, sheet, rows: Iterable[tuple[int, list]], next_row: int
    ) -> int:
        """Дописать строки [(row_index, [cell, ...])] в лист; пропуски заполняются пустыми.

        CPU-bound — звать из потока. Возвращает индекс следующей строки.
        """
        for row_index, cells in rows:
            for _ in range(row_index - next_row):
                sheet.append([])
            line: list[WriteOnlyCell | None] = []
            for cell in cells:
                line.extend([None] * (cell.col - len(line)))
                line.append(self._export_cell(sheet, cell))
            sheet.append(line)
            next_row = row_index + 1
        return next_row

    def _export_cell(self, sheet, cell) -> WriteOnlyCell:
        if cell.raw is not None and cell.raw.startswith("="):
            content = cell.raw
        elif isinstance(cell.value, dict):
            content = cell.value.get("e")
        elif cell.value is not None:
            content = cell.value
        else:
            content = cell.raw

        out = WriteOnlyCell(sheet, value=content)
        style = cell.style or {}
        if style.get("bg"):
            out.fill = PatternFill("solid", fgColor=style["bg"].lstrip("#"))
        if style.get("bold") or style.get("italic") or style.get("fg"):
            out.font = Font(
                bold=bool(style.get("bold")),
                italic=bool(style.get("italic")),
                color=style["fg"].lstrip("#") if style.get("fg") else None,
            )
        if style.get("halign") or style.get("wrap"):
            out.alignment = Alignment(
                horizontal=style.get("halign"), wrap_text=bool(style.get("wrap"))
            )
        if style.get("numFmt") in EXCEL_NUMBER_FORMATS:
            out.number_format = EXCEL_NUMBER_FORMATS[style["numFmt"]]
        return out

    @staticmethod
    def finalize_workbook(workbook: Workbook) -> BytesIO:
        """Сохранить книгу в буфер. CPU-bound — звать из потока."""
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer
